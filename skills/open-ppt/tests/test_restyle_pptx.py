"""Regression tests on native PPTX objects rather than screenshots."""

import hashlib
import json
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import restyle_pptx as helper


class RestyleTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / "source.pptx"
        self.plan_path = self.root / "plan.json"
        self.output = self.root / "mindray.pptx"
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = "项目现状与改进范围"
        slide.shapes.title.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0, 0, 255)
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.5), Inches(1.5), Inches(3), Inches(1))
        box.text = "保留原始内容：增长 20%，单位 mg/L。"
        box.text_frame.paragraphs[0].runs[0].hyperlink.address = "https://example.com/evidence"
        self.body_id = box.shape_id
        table = slide.shapes.add_table(2, 2, Inches(.5), Inches(3), Inches(3), Inches(1))
        table.table.cell(0, 0).text = "指标"
        table.table.cell(0, 1).text = "数值"
        table.table.cell(1, 0).text = "准确率"
        table.table.cell(1, 1).text = "99.2%"
        self.table_id = table.shape_id
        chart_data = CategoryChartData()
        chart_data.categories = ["基线", "当前"]
        chart_data.add_series("检测数量", [20, 30])
        chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(4), Inches(1.5), Inches(4), Inches(2.5), chart_data)
        self.chart_id = chart.shape_id
        group = slide.shapes.add_group_shape()
        group.shapes.add_textbox(Inches(4), Inches(4.2), Inches(3), Inches(.4)).text = "分组说明"
        self.group_id = group.shape_id
        image = slide.shapes.add_picture(str(helper.LOGO), Inches(8), Inches(0), width=Inches(1))
        self.image_id = image.shape_id
        slide.notes_slide.notes_text_frame.text = "必须保留的备注及数字 2026。"
        second = deck.slides.add_slide(deck.slide_layouts[5])
        second.shapes.title.text = "后续安排"
        second.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1)).text = "负责人和日期保持原样。"
        deck.save(self.source)
        self.source_hash = hashlib.sha256(self.source.read_bytes()).hexdigest()
        helper.inspect(self.source, self.plan_path)
        self.plan = json.loads(self.plan_path.read_text(encoding="utf-8"))

    def save_plan(self):
        self.plan_path.write_text(json.dumps(self.plan, ensure_ascii=False), encoding="utf-8")

    def test_native_content_tables_charts_notes_and_hyperlinks_survive(self):
        self.plan["slides"][0]["placements"] = [{"shape_id": self.chart_id, "preserve_style": False}]
        self.plan["slides"][1]["page_type"] = "closing"
        self.save_plan()
        helper.restyle(self.source, self.plan_path, self.output)
        deck = Presentation(self.output)
        self.assertEqual(len(deck.slides), 2)
        self.assertEqual((deck.slide_width, deck.slide_height), (960 * helper.EMU, 540 * helper.EMU))
        for number, slide in enumerate(deck.slides):
            current = {shape.shape_id: shape for shape in slide.shapes}
            for original in self.plan["slides"][number]["source_shapes"]:
                self.assertEqual(helper.text_values(current[original["shape_id"]]), original["text"])
        first = deck.slides[0]
        shapes = {shape.shape_id: shape for shape in first.shapes}
        self.assertEqual(shapes[self.body_id].text_frame.paragraphs[0].runs[0].hyperlink.address, "https://example.com/evidence")
        self.assertEqual(shapes[self.table_id].table.cell(1, 1).text, "99.2%")
        border = shapes[self.table_id].table.cell(1, 1)._tc.find(".//" + qn("a:lnL"))
        self.assertEqual(border.find(".//" + qn("a:srgbClr")).get("val"), helper.GRAY)
        self.assertEqual(tuple(shapes[self.chart_id].chart.series[0].values), (20.0, 30.0))
        self.assertEqual(shapes[self.chart_id].chart.series[0].format.fill.fore_color.rgb, RGBColor.from_string(helper.RED))
        self.assertEqual(first.notes_slide.notes_text_frame.text, "必须保留的备注及数字 2026。")
        self.assertEqual(first.shapes.title.text_frame.paragraphs[0].runs[0].font.color.rgb, RGBColor.from_string(helper.RED))
        self.assertEqual(str(deck.slides[1].background.fill.fore_color.rgb), helper.RED)
        self.assertTrue(any(s.name == "open-ppt Mindray logo" for s in first.shapes))
        self.assertFalse(any(s.name == "open-ppt Mindray logo" for s in deck.slides[1].shapes))
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(), self.source_hash)
        report = json.loads(self.output.with_suffix(".restyle-report.json").read_text(encoding="utf-8"))
        self.assertTrue(report["source_text_unchanged"])
        self.assertEqual(report["visual_validation"], "not_performed")
        with zipfile.ZipFile(self.source) as before, zipfile.ZipFile(self.output) as after:
            for part in before.namelist():
                if part.startswith("ppt/embeddings/"):
                    self.assertEqual(before.read(part), after.read(part))

    def test_plan_repositions_native_table_and_preserves_grid_data(self):
        self.plan["slides"][0]["placements"] = [{"shape_id": self.table_id, "bounds": [48, 260, 360, 120], "font_size_pt": 12}]
        self.save_plan()
        helper.restyle(self.source, self.plan_path, self.output)
        table_shape = next(s for s in Presentation(self.output).slides[0].shapes if s.shape_id == self.table_id)
        self.assertAlmostEqual(table_shape.left / helper.EMU, 48)
        self.assertAlmostEqual(table_shape.width / helper.EMU, 360)
        self.assertAlmostEqual(sum(c.width for c in table_shape.table.columns) / helper.EMU, 360, places=2)
        self.assertEqual(table_shape.table.cell(1, 1).text, "99.2%")

    def test_only_explicit_text_free_decoration_can_be_removed(self):
        self.plan["slides"][0]["decorative_shape_ids"] = [self.body_id]
        self.plan["slides"][0]["decoration_reason"] = "not permitted to remove business text"
        self.save_plan()
        with self.assertRaisesRegex(ValueError, "text-free"):
            helper.restyle(self.source, self.plan_path, self.output)
        self.assertFalse(self.output.exists())
        self.plan["slides"][0]["decorative_shape_ids"] = [self.image_id]
        self.plan["slides"][0]["decoration_reason"] = "Replace reviewed old logo with fixed Mindray logo"
        self.save_plan()
        helper.restyle(self.source, self.plan_path, self.output)
        report = json.loads(self.output.with_suffix(".restyle-report.json").read_text(encoding="utf-8"))
        self.assertEqual(report["removed_decorations"][0]["shape_id"], self.image_id)

    def test_stale_plan_unknown_shape_and_image_stretch_fail_without_output(self):
        baseline = json.dumps(self.plan)
        for change in ("hash", "inventory", "shape", "image", "master", "pages"):
            with self.subTest(change=change):
                self.plan = json.loads(baseline)
                if change == "hash":
                    self.plan["source_sha256"] = "stale"
                elif change == "inventory":
                    self.plan["slides"][0]["source_shapes"][0]["text"] = ["modified"]
                elif change == "shape":
                    self.plan["slides"][0]["placements"] = [{"shape_id": 9999}]
                elif change == "image":
                    self.plan["slides"][0]["placements"] = [{"shape_id": self.image_id, "bounds": [48, 100, 100, 100]}]
                elif change == "master":
                    self.plan["slides"][0]["hide_master_shapes"] = True
                else:
                    self.plan["slides"].pop()
                self.save_plan()
                with self.assertRaises(ValueError):
                    helper.restyle(self.source, self.plan_path, self.output)
                self.assertFalse(self.output.exists())
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(), self.source_hash)

    def test_existing_results_are_not_overwritten(self):
        self.output.write_bytes(b"previous output")
        with self.assertRaisesRegex(ValueError, "already exists"):
            helper.restyle(self.source, self.plan_path, self.output)
        self.assertEqual(self.output.read_bytes(), b"previous output")


if __name__ == "__main__":
    unittest.main()
