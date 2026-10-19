#!/usr/bin/env python3
"""Inspect a PPTX and apply a reviewed per-slide Mindray Red layout plan.

The agent chooses page roles, title objects, and any individual placements.
Native text, tables, charts, media, notes, and relationships are retained.
Requires python-pptx; no translation service or PPTD roundtrip is used.
"""

import argparse
import hashlib
import json
import math
from pathlib import Path
import tempfile
import xml.etree.ElementTree as ET
import zipfile

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.oxml.ns import qn
    from pptx.util import Pt
except ImportError:
    raise SystemExit("Mode 3 requires python-pptx: python -m pip install python-pptx")

from translate_pptx import read_package, write_json


EMU = 12700
WIDTH, HEIGHT = 960, 540
RED, INK, GRAY = "C8102E", "1A1A1A", "D9D9D9"
LOGO = Path(__file__).resolve().parents[1] / "reference/design_system/healthcare/mr-red-tech/assets/mindray-lockup-wide.png"


def text_values(shape):
    return [node.text or "" for node in shape._element.iter(qn("a:t"))]


def chart_values(shape):
    if not shape.has_chart:
        return None
    return [(node.tag, node.text) for node in shape.chart._chartSpace.iter()
            if node.tag in {qn("c:v"), qn("c:f"), qn("a:t")}]


def shape_record(shape):
    kind = ("group" if shape.shape_type == MSO_SHAPE_TYPE.GROUP else
            "picture" if shape.shape_type == MSO_SHAPE_TYPE.PICTURE else
            "chart" if shape.has_chart else "table" if shape.has_table else
            "text" if shape.has_text_frame else "shape")
    return {"shape_id": shape.shape_id, "name": shape.name, "kind": kind,
            "text": text_values(shape),
            "bounds": [round(v / EMU, 3) for v in (shape.left, shape.top, shape.width, shape.height)]}


def inspect(source, output):
    parts = read_package(source)
    deck = Presentation(source)
    if output.exists():
        raise ValueError("inspection file already exists; choose a new path")
    media_dir = output.parent / "source-media"
    media = []
    for part, data in parts.items():
        if part.startswith("ppt/media/") and not part.endswith("/"):
            name = hashlib.sha256(data).hexdigest()[:16] + Path(part).suffix
            destination = media_dir / name
            if not destination.exists():
                media_dir.mkdir(parents=True, exist_ok=True)
                destination.write_bytes(data)
            media.append({"part": part, "file": f"source-media/{name}"})
    slides = []
    for number, slide in enumerate(deck.slides, 1):
        title = slide.shapes.title
        records = [shape_record(shape) for shape in slide.shapes]
        slides.append({"slide": number, "page_type": "content",
                       "title_shape_id": title.shape_id if title is not None else None,
                       "source_shapes": records, "placements": [], "decorative_shape_ids": [],
                       "hide_master_shapes": False, "hide_master_shapes_reason": "",
                       "inherited_shape_candidates": {
                           "layout": [shape_record(s) for s in slide.slide_layout.shapes],
                           "master": [shape_record(s) for s in slide.slide_layout.slide_master.shapes]}})
    document = {"schema_version": 1, "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                "source_name": source.name, "slide_count": len(deck.slides),
                "source_size": [deck.slide_width / EMU, deck.slide_height / EMU],
                "output_size": [WIDTH, HEIGHT], "slides": slides, "media": media}
    write_json(output, document)
    return {"plan": str(output), "slides": len(slides), "media": len(media)}


def font_style(font, size, color):
    font.name = "Segoe UI"
    font.size = Pt(size)
    font.color.rgb = RGBColor.from_string(color)
    properties = font._element
    east_asian = properties.find(qn("a:ea"))
    if east_asian is None:
        east_asian = OxmlElement("a:ea")
        latin = properties.find(qn("a:latin"))
        if latin is not None:
            latin.addnext(east_asian)
        else:
            properties.insert(0, east_asian)
    east_asian.set("typeface", "Microsoft YaHei")


def frame_style(frame, size, color):
    for paragraph in frame.paragraphs:
        font_style(paragraph.font, size, color)
        for run in paragraph.runs:
            font_style(run.font, size, color)


def style_shape(shape, size, color, warnings, location, preserve_style=False):
    if preserve_style or shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            style_shape(child, size, color, warnings, location)
        return
    if shape.has_chart:
        # Native chart data and cached labels remain untouched. Recolor only
        # when the plan explicitly elects to style an ordinary data chart.
        chart = shape.chart
        palette = [RED, "BFBFBF", "8F0C20", INK, "595959", "D9D9D9"]
        for index, series in enumerate(chart.series):
            fill = RGBColor.from_string(palette[index % len(palette)])
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = fill
            series.format.line.color.rgb = fill
        warnings.append({"slide": location, "kind": "chart_review",
                         "detail": "Series palette updated; inspect markers, point overrides, labels, and semantic color coding."})
        return
    if shape.has_table:
        for row_number, row in enumerate(shape.table.rows):
            for cell in row.cells:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(RED if row_number == 0 else "FFFFFF")
                frame_style(cell.text_frame, size, "FFFFFF" if row_number == 0 else INK)
                properties = cell._tc.get_or_add_tcPr()
                for index, side in enumerate(("lnL", "lnR", "lnT", "lnB")):
                    existing = properties.find(qn("a:" + side))
                    if existing is not None:
                        properties.remove(existing)
                    line = OxmlElement("a:" + side)
                    line.set("w", "9525")
                    fill = OxmlElement("a:solidFill")
                    rgb = OxmlElement("a:srgbClr")
                    rgb.set("val", GRAY)
                    fill.append(rgb)
                    line.append(fill)
                    properties.insert(index, line)
        return
    if shape.has_text_frame:
        frame_style(shape.text_frame, size, color)
    if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
        geometry = shape._element.find(".//" + qn("a:prstGeom"))
        if geometry is not None and geometry.get("prst") == "roundRect":
            geometry.set("prst", "rect")
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string("8F0C20" if color == "FFFFFF" else "F5F5F5")
        shape.line.color.rgb = RGBColor.from_string(GRAY)
        shape.line.width = Pt(0.75)
    for effect in list(shape._element.iter(qn("a:effectLst"))):
        effect.getparent().remove(effect)


def place(shape, bounds):
    if not isinstance(bounds, list) or len(bounds) != 4 or any(
            isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v) for v in bounds):
        raise ValueError("bounds must contain four finite numbers in the 960 x 540 canvas")
    x, y, width, height = bounds
    if x < 0 or y < 0 or width < 0 or height < 0 or width + height == 0 or x + width > WIDTH + .01 or y + height > HEIGHT + .01:
        raise ValueError(f"placement outside output canvas: {bounds}")
    old_width, old_height = shape.width, shape.height
    if shape.shape_type in {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.GROUP} and old_width and old_height:
        if width <= 0 or height <= 0 or abs((width / height) / (old_width / old_height) - 1) > .01:
            raise ValueError("picture/group placements must preserve their original aspect ratio")
    shape.left, shape.top = round(x * EMU), round(y * EMU)
    new_width, new_height = round(width * EMU), round(height * EMU)
    shape.width, shape.height = new_width, new_height
    if shape.has_table and old_width and old_height:
        for column in shape.table.columns:
            column.width = round(column.width * new_width / old_width)
        for row in shape.table.rows:
            row.height = round(row.height * new_height / old_height)


def add_text(slide, name, text, bounds, size, color):
    box = slide.shapes.add_textbox(*[round(v * EMU) for v in bounds])
    box.name = name
    box.text_frame.text = text
    frame_style(box.text_frame, size, color)
    return box


def add_brand(slide, number, dark):
    if dark:
        return
    add_text(slide, "open-ppt Mindray marker", "❯", [48, 28, 20, 40], 24, RED)
    add_text(slide, "open-ppt page number", str(number), [762, 503, 36, 24], 9, "8C8C8C")
    logo = slide.shapes.add_picture(str(LOGO), round(812 * EMU), round(503 * EMU), width=round(100 * EMU))
    logo.name = "open-ppt Mindray logo"
    # A native line separates the header, with no decorative panel.
    line = slide.shapes.add_connector(1, round(48 * EMU), round(90 * EMU), round(912 * EMU), round(90 * EMU))
    line.name = "open-ppt header rule"
    line.line.color.rgb = RGBColor.from_string(GRAY)
    line.line.width = Pt(0.75)


def restyle(source, plan_path, output, keep_source_fonts=False):
    report_path = output.with_suffix(".restyle-report.json")
    if output.suffix.lower() != ".pptx" or len({p.resolve() for p in (source, plan_path, output, report_path)}) != 4:
        raise ValueError("use distinct source, plan, .pptx output, and report paths")
    if output.exists() or report_path.exists():
        raise ValueError("output/report already exists; choose a new output filename")
    parts = read_package(source)
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    if plan.get("schema_version") != 1 or plan.get("source_sha256") != hashlib.sha256(source.read_bytes()).hexdigest():
        raise ValueError("plan does not match the unchanged source PPTX")
    deck = Presentation(source)
    entries = plan.get("slides")
    if not isinstance(entries, list) or [s.get("slide") for s in entries] != list(range(1, len(deck.slides) + 1)):
        raise ValueError("plan must contain every source slide once, in presentation order")
    if not LOGO.is_file():
        raise ValueError(f"missing bundled Mindray logo: {LOGO}")
    warnings, removed, source_snapshots = [], [], []
    for number, (slide, entry) in enumerate(zip(deck.slides, entries), 1):
        page_type = entry.get("page_type")
        if page_type not in {"cover", "content", "divider", "closing"}:
            raise ValueError(f"slide {number}: invalid page_type")
        dark = page_type in {"cover", "closing"}
        original = {shape.shape_id: shape for shape in slide.shapes}
        before = {key: text_values(shape) for key, shape in original.items()}
        chart_before = {key: chart_values(shape) for key, shape in original.items() if shape.has_chart}
        if entry.get("source_shapes") != [shape_record(shape) for shape in slide.shapes]:
            raise ValueError(f"slide {number}: source shape inventory was modified")
        title_id = entry.get("title_shape_id")
        if title_id is not None and (title_id not in original or not original[title_id].has_text_frame):
            raise ValueError(f"slide {number}: title_shape_id must identify an existing text shape")
        decorative = entry.get("decorative_shape_ids", [])
        if not isinstance(decorative, list) or len(set(decorative)) != len(decorative):
            raise ValueError("decorative_shape_ids must be a list of unique shape IDs")
        for shape_id in decorative:
            shape = original.get(shape_id)
            if shape is None or shape_id == title_id or before[shape_id] or shape.has_chart or shape.has_table or shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                raise ValueError("only explicitly identified text-free decoration/old logo shapes may be removed")
            if not entry.get("decoration_reason"):
                raise ValueError("removing decorations requires a recorded decoration_reason")
            # Do not remove an animation/connector target and leave a dangling ID.
            for node in slide._element.iter():
                if node.get("spid") == str(shape_id) or (node.tag in {qn("a:stCxn"), qn("a:endCxn")} and node.get("id") == str(shape_id)):
                    raise ValueError("decoration is referenced by an animation/connector; preserve or edit it natively")
            shape._element.getparent().remove(shape._element)
            removed.append({"slide": number, "shape_id": shape_id, "reason": entry["decoration_reason"]})
        if entry.get("hide_master_shapes"):
            if not entry.get("hide_master_shapes_reason"):
                raise ValueError("hiding master shapes requires a recorded inspection reason")
            slide._element.set("showMasterSp", "0")
            warnings.append({"slide": number, "kind": "inherited_content_review",
                             "detail": entry["hide_master_shapes_reason"]})
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(RED if dark else "FFFFFF")
        items = entry.get("placements", [])
        if not isinstance(items, list):
            raise ValueError("placements must be a list")
        placements = {item["shape_id"]: item for item in items}
        if len(placements) != len(items) or not set(placements).issubset(set(original) - set(decorative)):
            raise ValueError("placement IDs must be unique source shape IDs")
        body = [shape for key, shape in original.items() if key != title_id and key not in decorative]
        # Fit existing content as a whole; the agent can override each object
        # for genuine re-layout using placements after inspecting the slide.
        region = [64, 324, 832, 150] if dark else [48, 110, 864, 370]
        if body:
            left = min(s.left for s in body) / EMU
            top = min(s.top for s in body) / EMU
            width = max(s.left + s.width for s in body) / EMU - left
            height = max(s.top + s.height for s in body) / EMU - top
            scale = min(region[2] / max(width, .001), region[3] / max(height, .001))
        for shape in list(slide.shapes):
            item = placements.get(shape.shape_id, {})
            is_title = shape.shape_id == title_id
            bounds = item.get("bounds")
            if bounds is None:
                bounds = ([64, 160, 832, 140] if dark else [72, 28, 840, 58]) if is_title else [
                    region[0] + (shape.left / EMU - left) * scale,
                    region[1] + (shape.top / EMU - top) * scale,
                    shape.width / EMU * scale, shape.height / EMU * scale]
            place(shape, bounds)
            size = item.get("font_size_pt", 36 if dark and is_title else 26 if is_title else 12 if shape.has_table else 14)
            if isinstance(size, bool) or not isinstance(size, (int, float)) or not math.isfinite(size) or not 8 <= size <= 96:
                raise ValueError("font_size_pt must be between 8 and 96 points; re-layout rather than hiding text in tiny type")
            preserve = item.get("preserve_style", shape.has_chart)
            style_shape(shape, size, "FFFFFF" if dark else RED if is_title else INK, warnings, number, preserve)
            if shape.has_chart and preserve:
                warnings.append({"slide": number, "shape_id": shape.shape_id, "kind": "chart_style_preserved",
                                 "detail": "Native chart colors retained; review semantic palette before recoloring."})
        after = {shape.shape_id: text_values(shape) for shape in slide.shapes}
        if after != {key: value for key, value in before.items() if key not in decorative}:
            raise ValueError(f"slide {number}: source text changed during styling")
        source_snapshots.append(({key: value for key, value in before.items() if key not in decorative}, chart_before))
        if title_id is None:
            warnings.append({"slide": number, "kind": "missing_title", "detail": "No source title chosen; review the page header."})
        add_brand(slide, number, dark)
    deck.slide_width, deck.slide_height = WIDTH * EMU, HEIGHT * EMU
    if not keep_source_fonts:
        for node in list(deck._element):
            if node.tag == qn("p:embeddedFontLst"):
                deck._element.remove(node)
        for key in ("embedTrueTypeFonts", "saveSubsetFonts"):
            deck._element.attrib.pop(key, None)
        for rel in list(deck.part.rels.values()):
            if rel.reltype.endswith("/font"):
                deck.part.drop_rel(rel.rId)
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".restyle-", dir=output.parent) as temporary:
        staged = Path(temporary) / "result.pptx"
        deck.save(staged)
        check = Presentation(staged)
        if len(check.slides) != len(deck.slides):
            raise ValueError("slide count changed unexpectedly")
        for number, (slide, (texts, charts)) in enumerate(zip(check.slides, source_snapshots), 1):
            shapes = {shape.shape_id: shape for shape in slide.shapes}
            for key, values in texts.items():
                if key not in shapes or text_values(shapes[key]) != values:
                    raise ValueError(f"slide {number}: saved source text differs from original")
            for key, values in charts.items():
                if key not in shapes or chart_values(shapes[key]) != values:
                    raise ValueError(f"slide {number}: native chart data or labels changed")
        with zipfile.ZipFile(staged) as archive:
            if archive.testzip():
                raise ValueError("generated PPTX ZIP integrity failed")
            generated = {name: archive.read(name) for name in archive.namelist()}
        protected = [name for name in parts if name.startswith(("ppt/embeddings/", "ppt/notesSlides/", "ppt/media/")) and not name.endswith("/")]
        for part in protected:
            # Decoration removal may leave unused images. Business-image shapes
            # and embedded workbooks stay related and are checked below.
            if part.startswith("ppt/media/") and part not in generated and removed:
                continue
            if part.startswith("ppt/notesSlides/") and part.endswith((".xml", ".rels")) and part in generated:
                same = ET.canonicalize(parts[part].decode("utf-8"), rewrite_prefixes=True) == ET.canonicalize(
                    generated[part].decode("utf-8"), rewrite_prefixes=True)
            else:
                same = generated.get(part) == parts[part]
            if not same:
                raise ValueError(f"protected source content was not preserved: {part}")
        with output.open("xb") as handle:
            handle.write(staged.read_bytes())
    report = {"source": str(source.resolve()), "output": str(output.resolve()),
              "style": "mr-red-tech", "slide_count": len(deck.slides), "output_size": [WIDTH, HEIGHT],
              "source_text_unchanged": True, "native_chart_data_unchanged": True, "protected_parts_checked": protected,
              "removed_decorations": removed, "font_embedding_retained": keep_source_fonts,
              "visual_validation": "not_performed", "warnings": warnings}
    write_json(report_path, report)
    return {"output": str(output), "report": str(report_path), "slides": len(deck.slides)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    get = commands.add_parser("inspect")
    get.add_argument("source", type=Path)
    get.add_argument("--output", type=Path, required=True)
    put = commands.add_parser("apply")
    put.add_argument("source", type=Path)
    put.add_argument("--plan", type=Path, required=True)
    put.add_argument("--output", type=Path, required=True)
    put.add_argument("--keep-source-fonts", action="store_true")
    args = parser.parse_args()
    try:
        result = inspect(args.source, args.output) if args.command == "inspect" else restyle(
            args.source, args.plan, args.output, args.keep_source_fonts)
    except (OSError, ValueError, KeyError, TypeError, zipfile.BadZipFile) as error:
        parser.exit(1, f"Mindray restyle: {error}\n")
    print(json.dumps(result, ensure_ascii=True))


if __name__ == "__main__":
    main()
